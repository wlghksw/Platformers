import os
import shutil
import glob
import json
import time
import re
import base64
import markdown
from datetime import datetime
from typing import List
from io import BytesIO
from PIL import Image
from mimetypes import guess_type
from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
import anthropic

# Document parsing libraries
from pptx import Presentation
from docx import Document
import pdfplumber  # Replaced fitz due to DLL issues
from fpdf import FPDF

app = FastAPI()

# Enable CORS for Next.js frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Base directories
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT_DIR = os.path.join(BASE_DIR, "data", "input")
OUTPUT_DIR = os.path.join(BASE_DIR, "data", "output")

# Ensure directories exist
os.makedirs(INPUT_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

app.mount("/output", StaticFiles(directory=OUTPUT_DIR), name="output")

HISTORY_FILE = os.path.join(BASE_DIR, "data", "history.json")
if not os.path.exists(HISTORY_FILE):
    with open(HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump([], f)

class ProcessRequest(BaseModel):
    analysis_link: str = ""
    card_link: str = ""
    reference_links: str = ""
    extra_hashtags: str = ""
    api_key: str

class RenderRequest(BaseModel):
    json_data: dict
    used_images: list
    layout_images: dict  # mapping: paragraph index -> image filename (e.g. {"1": "abc.png"})
    base_filename: str
    start_time: float

def encode_image(image_path):
    max_size = 3 * 1024 * 1024  # 최대 3MB 수준으로 압축 (Claude 제한 5MB 대비 안전치)
    with Image.open(image_path) as img:
        # P나 RGBA 모드일 경우 JPEG 저장을 위해 RGB로 변환
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        
        # 해상도가 너무 클 경우 2000px로 리사이징
        max_dim = 2000
        if img.width > max_dim or img.height > max_dim:
            img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)
        
        quality = 85
        buffer = BytesIO()
        while True:
            buffer.seek(0)
            buffer.truncate()
            # JPEG로 퀄리티를 조절하여 저장 시도
            img.save(buffer, format="JPEG", quality=quality)
            
            # 크기가 3MB 이하이거나, 더 이상 압축하기 어려울 경우 탈출
            if buffer.tell() <= max_size or quality <= 20:
                break
            quality -= 15
            
        buffer.seek(0)
        return base64.b64encode(buffer.read()).decode('utf-8'), "image/jpeg"

def extract_text_from_pptx(filepath):
    prs = Presentation(filepath)
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_runs.append(shape.text)
    return "\n".join(text_runs)

def extract_text_from_docx(filepath):
    doc = Document(filepath)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_from_pdf(filepath):
    text = ""
    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ""
    except Exception as e:
        print(f"pdfplumber error: {e}")
    return text

@app.post("/clear_input")
async def clear_input():
    for f in glob.glob(os.path.join(INPUT_DIR, "*")):
        try:
            if os.path.isfile(f):
                os.unlink(f)
        except Exception as e:
            print(f"Delete error: {e}")
    return {"status": "success"}

@app.post("/upload")
async def upload_files(files: List[UploadFile] = File(...)):
    # 이전 세션 파일 모두 지우기 (한 번에 분석할 파일만 남김)
    for f in glob.glob(os.path.join(INPUT_DIR, "*")):
        try:
            os.remove(f)
        except Exception:
            pass

    saved_files = []
    for file in files:
        file_path = os.path.join(INPUT_DIR, file.filename)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        saved_files.append(file.filename)
    return {"status": "success", "files": saved_files}

import uuid
@app.post("/upload_single_image")
async def upload_single_image(file: UploadFile = File(...), base_filename: str = Form(None)):
    if not base_filename:
        base_filename = f"blog_{int(time.time())}"
        
    ext = os.path.splitext(file.filename)[1].lower()
    if not ext:
        ext = ".png"
        
    img_filename = f"{base_filename}_inline_{uuid.uuid4().hex[:6]}{ext}"
    out_path = os.path.join(OUTPUT_DIR, img_filename)
    
    with open(out_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    return {"status": "success", "filename": img_filename}

@app.post("/generate")
async def generate_blog(request: ProcessRequest):
    start_time = time.time()
    file_count = len(glob.glob(os.path.join(OUTPUT_DIR, "blog_*.html"))) + 1
    base_filename = f"blog_{file_count}"

    import requests
    from bs4 import BeautifulSoup
    
    scraped_images = []
    fetched_title = "제목 미상"
    all_text = ""
    
    if request.analysis_link and request.analysis_link.strip().startswith('http'):
        try:
            headers = {'User-Agent': 'Mozilla/5.0'}
            res = requests.get(request.analysis_link.strip(), headers=headers, timeout=10)
            if res.status_code == 200:
                soup = BeautifulSoup(res.text, 'html.parser')
                
                h1_tag = soup.find('h1', class_='entry-title')
                if h1_tag:
                    fetched_title = h1_tag.get_text(separator=" ", strip=True)
                
                all_text += f"프로젝트/강좌명: {fetched_title}\n\n"

                # 네이버 블로그 스마트에디터 iframe 대응
                if 'blog.naver.com' in request.analysis_link and 'PostView' not in request.analysis_link:
                    iframe = soup.find('iframe', id='mainFrame')
                    if iframe and iframe.get('src'):
                        real_url = "https://blog.naver.com" + iframe['src']
                        res = requests.get(real_url, headers=headers, timeout=10)
                        soup = BeautifulSoup(res.text, 'html.parser')
                
                # 이미지 추출 로직 (본문 내 의미있는 사진 필터링)
                import uuid
                p_tags = soup.find_all('p', style=lambda value: value and 'text-align: center' in value.lower())
                img_tags = []
                for p in p_tags:
                    img_tags.extend(p.find_all('img'))
                for img in img_tags:
                    if len(scraped_images) >= 15:
                        break
                    src = img.get('data-lazy-src') or img.get('src')
                    if not src or not src.startswith('http'):
                        continue
                    img_class = img.get('class', [])
                    if isinstance(img_class, list):
                        img_class_str = " ".join(img_class).lower()
                    else:
                        img_class_str = str(img_class).lower()
                    
                    try:
                        img_width = int(img.get('width', 0))
                    except:
                        img_width = 0

                    is_valid_image = False
                    # 1. 네이버 블로그 도메인/속성
                    if 'postfiles.pstatic.net' in src or 'blogfiles.naver.net' in src or 'se-image-resource' in img_class_str:
                        is_valid_image = True
                    # 2. 워드프레스 및 범용 HTML 속성
                    elif 'wp-image' in img_class_str or 'size-full' in img_class_str or 'aligncenter' in img_class_str:
                        is_valid_image = True
                    # 3. 크기가 명시된 사진
                    elif img_width >= 250:
                        is_valid_image = True

                    if is_valid_image:
                        try:
                            img_res = requests.get(src, headers=headers, timeout=5)
                            if img_res.status_code == 200 and len(img_res.content) > 15000: # 15KB 이상 진짜 사진만
                                ext = os.path.splitext(src.split('?')[0])[1].lower()
                                if ext not in ['.jpg', '.jpeg', '.png', '.gif']: 
                                    ext = ".jpg"
                                img_filename = f"{base_filename}_scraped_{uuid.uuid4().hex[:6]}{ext}"
                                out_path = os.path.join(OUTPUT_DIR, img_filename)
                                with open(out_path, "wb") as f:
                                    f.write(img_res.content)
                                scraped_images.append(img_filename)
                        except Exception:
                            pass
                # 불필요한 스크립트/스타일 제거
                for script in soup(["script", "style", "nav", "footer", "iframe"]):
                    script.extract()
                extracted_text = soup.get_text(separator=' ', strip=True)
                all_text += f"\n\n[웹사이트 크롤링 원문 텍스트 ({request.analysis_link})]\n{extracted_text[:4000]}\n\n"
            else:
                all_text += f"\n\n(웹사이트 크롤링 실패 - 상태코드: {res.status_code})\n\n"
        except Exception as e:
            all_text += f"\n\n(웹사이트 크롤링 오류: {str(e)})\n\n"
    image_contents = []
    used_images = []

    if hasattr(request, 'card_link') and request.card_link and request.card_link.strip().startswith("http"):
        _card_link = request.card_link.strip()
        _card_filename = f"{base_filename}_card_screenshot.png"
        _card_save_path = os.path.join(OUTPUT_DIR, _card_filename)

        def _capture_card():
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                pg = browser.new_page()
                pg.goto(_card_link, timeout=15000)
                card = pg.wait_for_selector('.wpgb-card-inner', timeout=8000)
                if card:
                    card.screenshot(path=_card_save_path)
                browser.close()

        try:
            import concurrent.futures
            with concurrent.futures.ThreadPoolExecutor() as pool:
                future = pool.submit(_capture_card)
                future.result(timeout=30)
            if os.path.exists(_card_save_path):
                used_images.append(_card_filename)
                print(f"[Card Capture] 성공: {_card_filename}")
            else:
                print("[Card Capture] 파일 저장 실패")
        except Exception as e:
            print(f"[Card Capture] 오류: {str(e)}")
            all_text += f"\n\n(카드 링크 캡처 오류: {str(e)})\n\n"


    client = anthropic.Anthropic(api_key=request.api_key)
    
    image_instruction = ""

    text_prompt = f"""
# 지시어: 아래 정보를 바탕으로 네이버 블로그 스마트에디터 완벽 복제용 JSON 데이터를 생성해줘.

**[매우 중요한 기본 규칙 - 이모지 절대 금지 & 이미지 문구 절대 금지]**
- 절대로 어떠한 형태의 이모지나 이모티콘(🔥, 😊, 🚀, 👍 등)도 사용하지 마세요.
- 절대로 본문(body_paragraphs) 내에 `[IMAGE_X]` 와 같은 이미지 배치 기호를 넣지 마세요. 사용자가 알아서 배치합니다.
- 외부 사이트 크롤링 텍스트에 포함된 쓸모없는 네이버 블로그 문구(예: "존재하지 않는 이미지입니다.", "AI 활용 설정", "사진 설명을 입력하세요.")를 절대로 결과물에 포함하지 마세요. 철저히 지우고 핵심 내용만 추출하세요.

## [입력 데이터]
- 제목 키워드: {fetched_title}
- 참고 링크: {request.reference_links}
- 분석할 원본 링크: {request.analysis_link}
- 필수 해시태그: #에듀올랩 #크레용스쿨 #홈스쿨 #이러닝강좌 {request.extra_hashtags}

## [강력 지시사항: 스크래핑된 외부 문서 핵심 요약 및 엄격한 규칙]
하단 **[참고할 첨부파일 기반 추출 내용]** 섹션에 사용자가 제공한 **[웹사이트 크롤링 원문 텍스트]**가 포함되어 있다면, **반드시 그 내용을 1순위로 정독하고 가장 중요한 핵심 정보만 요약하여 간결하게** 작성하세요. 불필요하게 살을 붙이지 말고 가독성이 좋도록 짧고 명확한 문장을 사용하세요.

**[본문 작성 필수 규칙]**
1. 강의의 핵심 소재나 도구가 있다면 이에 대한 간단한 설명 문구를 추가하세요.
2. 강의를 통해서 배울 수 있는 것이 무엇인지에 대한 확실한 설명을 포함하세요.
3. 단, **위 1, 2번 규칙에 대한 내용은 반드시 '웹사이트 크롤링 원문 텍스트'에서만 근거를 찾아 작성해야 하며, 원문에 해당 내용이 존재하지 않는다면 절대로 임의로 지어내어 추가하지 마세요.**
4. **`main_title` 필드는 반드시 샘플 포맷 그대로 작성해야 합니다. 특히 앞부분의 'Home스쿨'은 절대로 '홈스쿨'이나 다른 한글로 바꾸지 마세요. 영문 'Home'과 한글 '스쿨'의 조합 형태 그대로 유지해야 합니다.**

## [출력 형식 (오직 JSON 데이터만 출력할 것)]
반드시 아래 JSON 포맷에 맞추어 코드 블록(```json) 안에 응답하세요. 다른 설명은 일절 생략하세요.

```json
{{
  "category": "사회/경제, 전문강좌 (혹은 성격에 맞는 분야)",
  "main_title": "Home스쿨 | {fetched_title} -크레용스쿨 이러닝 강좌 소개",
  "quote_hook": "부모님의 시선을 끄는 캐치프레이즈 인용구를 작성하세요.\\n마치 세계의 이야기와 랜드마크를 재미있게~ 처럼 작성합니다.",
  "body_paragraphs": [
    "{fetched_title} 는 초등학생을 위한 강좌입니다. (도입 첫 문장 등 자연스러운 시작)",
    "학생들은 자연스럽게 강좌의 핵심 내용과 도구의 사용법을 학습 가능합니다.",
    "본문 배열(body_paragraphs)은 무조건 최대 7문단(요소 7개) 이하로 구성하세요. 핵심 위주의 간결한 문장이어야 합니다. 절대로 이미지 마커는 넣지 마세요."
  ],
  "highlight_keywords": [
    "문화, 역사", "소근육 발달", "이해도가 더 높아집니다", "공간 지각 능력과 창의력"
  ],
  "youtube_link_title": "유튜브에서 즐기기",
  "youtube_link": "입력받은 참고링크 중 유튜브 형태의 링크 (없으면 빈칸)",
  "homepage_link_title": "만들면서 배우는 홈페이지 - 크레용스쿨",
  "homepage_link": "입력받은 참고링크 중 홈페이지 링크 (없으면 빈칸)",
  "outro_text": "소개 드리는 강좌들은\\n엄마도 손쉽게 교육할 수 있도록\\n제작한 홈스쿨 이러닝 강좌들입니다!\\n\\n크레용스쿨에서 제작/활용하게 될\\n교육 영상 콘텐츠를 기대해 주세요.\\n\\n앞으로 좋은 콘텐츠 제작에 앞장설\\n에듀올랩에\\n많은 기대와 관심 부탁드려요.",
  "hashtags": "해시태그 모음들"
}}
```
- body_paragraphs 내부에서 시선을 끌만한(주황색/볼드 적용 유도할) 핵심 문구나 단어를 뽑아서 highlight_keywords 배열에 넣어주세요.
- quote_hook, outro_text 내에서 줄바꿈이 필요하다면 반드시 \\n 으로 처리해 주세요.

{image_instruction}

[참고할 첨부파일 기반 추출 내용]
{all_text}
"""

    messages_content = []
    if image_contents:
        messages_content.extend(image_contents)
    messages_content.append({"type": "text", "text": text_prompt})

    try:
        response = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=4000,
            messages=[{"role": "user", "content": messages_content}]
        )
        blog_content = response.content[0].text
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    # Parse JSON
    import re
    import json
    
    json_data = {}
    json_match = re.search(r'```(?:json)?\s*(\{.*?\})\s*```', blog_content, re.DOTALL)
    if json_match:
        try:
            json_data = json.loads(json_match.group(1))
        except Exception as e:
            print("JSON parsing error:", e)
    else:
        print("Regex failed to find JSON format.")

    # main_title 강제 덮어씌우기 (AI가 번역/변형하지 못하도록)
    json_data["main_title"] = f"Home스쿨 | {fetched_title} -크레용스쿨 이러닝 강좌 소개"
    # 본문 h1용 원본 강좌명 별도 저장
    json_data["original_title"] = fetched_title
        
    return {"status": "success", "json_data": json_data, "used_images": used_images, "scraped_images": scraped_images, "base_filename": base_filename, "start_time": start_time}

@app.post("/render_html")
async def render_html(request: RenderRequest):
    import json
    json_data = request.json_data
    used_images = request.used_images
    layout_images = request.layout_images
    base_filename = request.base_filename
    start_time = request.start_time

    html_path = os.path.join(OUTPUT_DIR, f"{base_filename}.html")
    txt_path = os.path.join(OUTPUT_DIR, f"{base_filename}.txt")
    
    # Render Body Paragraphs with Highlights and User-placed Layout Images
    highlighted_paragraphs = []
    keywords = json_data.get('highlight_keywords', [])
    body_html_parts = []
    
    for idx, p_text in enumerate(json_data.get('body_paragraphs', [])):
        p = p_text
        for kw in keywords:
            if kw and kw in p:
                p = p.replace(kw, f'<span style="color:#ff8c00; font-weight:bold;">{kw}</span>')
        
        # Add paragraph HTML
        body_html_parts.append(f"<div style='text-align: center; color: #555; line-height: 1.8; font-size: 16px; font-weight: 500; font-family: \"Malgun Gothic\", sans-serif; margin-bottom: 25px;'>{p}</div>")
        
        # Check if user placed an image right after this paragraph (idx)
        str_idx = str(idx)
        if str_idx in layout_images and layout_images[str_idx]:
            img_name = layout_images[str_idx]
            body_html_parts.append(f"<div style='text-align: center; margin: 35px 0;'><img src='./{img_name}' style='max-width: 100%; border-radius: 12px; box-shadow: 0 4px 10px rgba(0,0,0,0.05);'/></div>")
    
    body_html = "\n".join(body_html_parts)

    quote_hook = json_data.get('quote_hook', '').replace('\\n', '<br>').replace('\n', '<br>')
    outro_text = json_data.get('outro_text', '').replace('\\n', '<br>').replace('\n', '<br>').replace('에듀올랩', '<span style="color: #ff9900; font-size: 24px; font-weight: bold; font-style: italic;">에듀올랩</span>')
    
    # Generate Link Cards
    link_cards_html = ""
    if json_data.get('homepage_link'):
        link_cards_html += f'''
        <div style="border: 1px solid #e5e5e5; border-radius: 8px; padding: 20px; text-align: left; margin: 20px auto; max-width: 600px;">
            <a href="{json_data.get('homepage_link')}" style="text-decoration: none; color: #333; display: block;">
                <div style="font-weight: bold; margin-bottom: 8px;">{json_data.get('homepage_link_title', '링크 바로가기')}</div>
                <div style="font-size: 13px; color: #00c73c;">{json_data.get('homepage_link')}</div>
            </a>
        </div>
        '''
    if json_data.get('youtube_link'):
        link_cards_html += f'''
        <div style="font-weight: bold; margin: 40px 0 10px 0; color:#333;">{json_data.get('youtube_link_title', '유튜브에서 즐기기')}</div>
        <div style="border: 1px solid #e5e5e5; border-radius: 8px; padding: 20px; text-align: left; margin: 0 auto; max-width: 600px;">
            <a href="{json_data.get('youtube_link')}" style="text-decoration: none; color: #333; display: block;">
                <div style="font-weight: bold; margin-bottom: 8px;">유튜브 영상 보기</div>
                <div style="font-size: 13px; color: #00c73c;">youtube.com</div>
            </a>
        </div>
        '''

    card_screenshot_path = ""
    for used in used_images:
        if used.endswith("_card_screenshot.png"):
            card_screenshot_path = used
            break

    card_image_html = f'<div style="text-align: center; margin: 40px 0;"><img src="http://localhost:8000/output/{card_screenshot_path}" style="max-width:100%; border-radius:12px; box-shadow: 0 4px 15px rgba(0,0,0,0.1); width:640px;"/></div>' if card_screenshot_path else ''

    # Assemble Full Naver Clone HTML Envelope
    html_template = f"""
    <!DOCTYPE html>
    <html lang="ko">
    <head><meta charset="UTF-8"><title>블로그 결과 - {base_filename}</title></head>
    <body style="margin: 0; padding: 0; background-color: #f9f9f9;">
    <div style="background-color: white; font-family: 'Apple SD Gothic Neo', 'Malgun Gothic', 'Dotum', sans-serif; max-width: 800px; margin: 40px auto; padding: 60px 40px; box-shadow: 0 4px 12px rgba(0,0,0,0.05);">
        
        <div style="font-size: 16px; font-weight: bold; margin-bottom: 30px; color: #333; padding: 15px; background-color: #e9ecef; border-radius: 8px; border-left: 4px solid #ff9900;">
            📝 블로그 작성 시 제목입력칸용 복사 내용: <br/>{json_data.get('main_title', '')}
        </div>

        <!-- header area -->
        <div style="text-align: center; margin-bottom: 40px;">
            <span style="color: #ff9900; font-weight: bold; font-size: 16px;">{json_data.get('category', '')}</span>
            <h1 style="font-size: 32px; font-weight: bold; margin: 20px 0; color: #222; word-break: keep-all; line-height: 1.4;">{json_data.get('original_title', json_data.get('main_title', ''))}</h1>
            <div style="margin: 40px auto 30px auto; border-top: 1px solid #777; width: 60%; position: relative;">
                <div style="position: absolute; top: -7px; left: 50%; width: 12px; height: 12px; background: white; border: 1px solid #555; transform: translateX(-50%) rotate(45deg);"></div>
            </div>
        </div>
        
        <!-- first image (if explicitly chosen, else default to first uploaded) -->
      
        <!-- captured card screenshot -->
        {card_image_html}

        <!-- quote hook -->
        <div style="text-align: center; margin: 70px 0;">
            <span style="font-size: 50px; color: #ccc; font-family: serif; display: block; height: 30px; line-height: 30px;">“</span>
            <p style="color: #ff9900; font-size: 20px; font-weight: bold; font-style: italic; line-height: 1.8; margin: 30px 0;">
                {quote_hook}
            </p>
            <span style="font-size: 50px; color: #ccc; font-family: serif; display: block; height: 30px; line-height: 30px;">”</span>
        </div>

        <!-- main body -->
        <div style="margin-bottom: 80px;">
            {body_html}
        </div>
        
        <!-- Links -->
        <div style="margin: 60px 0; text-align: center;">
            <div style="width: 1px; height: 40px; background-color: #aaa; margin: 0 auto 30px auto;"></div>
            {link_cards_html}
        </div>

        <!-- Outro -->
        <div style="text-align: center; font-size: 17px; color: #555; font-weight:500; line-height: 1.9; margin-top: 80px;">
            {outro_text}
            <br><br>
            <span style="color: #888; font-size: 15px;">{json_data.get('hashtags', '')}</span>
        </div>
        
    </div>
    </body>
    </html>
    """
    
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_template)

    # Clean TXT Dump (Just in Case)
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(json.dumps(json_data, ensure_ascii=False, indent=2))

    # Record history
    duration = time.time() - start_time
    try:
        with open(HISTORY_FILE, "r+", encoding="utf-8") as f:
            history = json.load(f)
            history.append({
                "date": datetime.now().strftime("%Y-%m-%d"),
                "duration": duration
            })
            f.seek(0)
            json.dump(history, f, indent=2)
    except Exception as e:
        print(f"History record error: {e}")

    return {"status": "success", "filename": f"{base_filename}.html"}

@app.get("/stats")
async def get_stats():
    total_files = len(glob.glob(os.path.join(OUTPUT_DIR, "blog_*.html")))
    today_str = datetime.now().strftime("%Y-%m-%d")
    
    today_count = 0
    total_duration = 0
    history_count = 0
    
    try:
        with open(HISTORY_FILE, "r", encoding="utf-8") as f:
            history = json.load(f)
            history_count = len(history)
            for item in history:
                if item.get("date") == today_str:
                    today_count += 1
                total_duration += item.get("duration", 0)
    except:
        pass
    
    avg_time = round(total_duration / history_count, 1) if history_count > 0 else 0
    
    return {
        "todayCount": today_count,
        "avgTime": avg_time,
        "totalCount": total_files
    }

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
